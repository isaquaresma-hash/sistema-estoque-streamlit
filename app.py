import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def gerar_excel_estilizado(df):
    buffer = io.BytesIO()
    
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Relatório')
        worksheet = writer.sheets['Relatório']
        
        # Estilo do Cabeçalho Verde Escuro
        header_fill = PatternFill(start_color="003319", end_color="003319", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        # Aplicar no Cabeçalho
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            
        worksheet.row_dimensions[1].height = 25

        # Estilo dos Dados
        data_font = Font(name="Calibri", size=11)
        data_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(bottom=Side(style='thin', color='E0E0E0'))

        for row_num in range(2, len(df) + 2):
            worksheet.row_dimensions[row_num].height = 20
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.font = data_font
                cell.alignment = data_alignment
                cell.border = thin_border

        # Largura Automática das Colunas
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            worksheet.column_dimensions[col_letter].width = max(max_len + 5, 14)

    buffer.seek(0)
    return buffer.getvalue()
